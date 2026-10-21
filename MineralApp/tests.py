"""
Pruebas automáticas de Mineral Manager.

Cada prueba parte de la faena de demostración (comando cargar_demo) y recorre los flujos
principales tal como lo haría un usuario. Ejecutar con:

    python manage.py test
"""
from datetime import timedelta
from io import BytesIO, StringIO

from django.core.management import call_command
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from .models import (
    ArticuloBodega, ArticuloPanol, Capacitacion, CapacitacionTrabajador, Maquinaria,
    MantenimientoMaquinaria, MovimientoArticulo, Panol, RetiroArticulo, Trabajador,
    TrabajoMaquinaria,
)


class BaseDemoTest(TestCase):
    """Carga la demo una vez por clase e inicia sesión con la cuenta demo en cada prueba."""

    @classmethod
    def setUpTestData(cls):
        call_command('cargar_demo', stdout=StringIO())

    def setUp(self):
        self.client.login(username='demo', password='demo1234')


class AccesoTests(TestCase):
    def setUp(self):
        call_command('cargar_demo', stdout=StringIO())

    def test_login_muestra_credenciales_demo(self):
        respuesta = self.client.get(reverse('login_signup'))
        self.assertContains(respuesta, 'demo1234')

    def test_login_correcto_lleva_al_panel(self):
        respuesta = self.client.post(reverse('login_signup'), {'username': 'demo', 'password': 'demo1234'})
        self.assertRedirects(respuesta, reverse('home'))

    def test_login_incorrecto_muestra_error(self):
        respuesta = self.client.post(reverse('login_signup'), {'username': 'demo', 'password': 'mala'})
        self.assertContains(respuesta, 'incorrectos')

    def test_paginas_privadas_requieren_sesion(self):
        for ruta in ['home', 'gestion_trabajadores', 'maquinaria', 'bodega', 'upload_csv']:
            respuesta = self.client.get(reverse(ruta))
            self.assertEqual(respuesta.status_code, 302, ruta)
        # La baja de trabajadores tampoco se puede hacer sin sesión
        self.client.post(reverse('eliminar_trabajador'), '{"rut": "12345678-9"}', content_type='application/json')
        self.assertTrue(Trabajador.objects.get(rut='12345678-9').activo)

    def test_cerrar_sesion(self):
        self.client.login(username='demo', password='demo1234')
        self.client.post(reverse('logout'))
        self.assertEqual(self.client.get(reverse('home')).status_code, 302)


class PaginasTests(BaseDemoTest):
    def test_todas_las_paginas_cargan(self):
        maquina = Maquinaria.objects.get(codigo_maquinaria='CAEX-001')
        rutas = [
            reverse('home'), reverse('gestion_trabajadores'), reverse('add_trabajador'),
            reverse('editar_trabajador', args=['12345678-9']), reverse('maquinaria'),
            reverse('ver_mantenimientos', args=[maquina.id]), reverse('ver_trabajos', args=[maquina.id]),
            reverse('pañol'), reverse('bodega'), reverse('retiro_articulo'), reverse('mover_articulo'),
            reverse('upload_csv'),
        ]
        for ruta in rutas:
            with self.subTest(ruta=ruta):
                self.assertEqual(self.client.get(ruta).status_code, 200)

    def test_informes_pdf(self):
        for ruta in ['descargar_informe_trabajadores', 'descargar_informe_bodega', 'descargar_informe_pañol']:
            respuesta = self.client.get(reverse(ruta))
            self.assertEqual(respuesta['Content-Type'], 'application/pdf', ruta)

    def test_filtros_y_busquedas(self):
        maquina = Maquinaria.objects.get(codigo_maquinaria='CAEX-001')
        self.assertContains(self.client.get(reverse('gestion_trabajadores'), {'certificacion': 'expirada'}), 'Juan Pérez')
        self.assertContains(self.client.get(reverse('ver_mantenimientos', args=[maquina.id]), {'search': 'Ana'}), 'aceite')
        self.assertContains(self.client.get(reverse('pañol'), {'cantidad_filtro': 'agotado'}), 'Esmeril')


class CargaDatosTests(BaseDemoTest):
    def test_demo_cargada_desde_excel(self):
        self.assertEqual(Trabajador.objects.count(), 12)
        self.assertEqual(Maquinaria.objects.count(), 8)
        self.assertEqual(CapacitacionTrabajador.objects.count(), 14)
        # Las horas de cada trabajo se calculan desde inicio y fin
        self.assertTrue(all(t.horas_trabajadas > 0 for t in TrabajoMaquinaria.objects.all()))

    def test_excel_con_error_no_guarda_nada(self):
        libro = Workbook()
        hoja = libro.active
        hoja.append(['Código', 'Nombre Maquinaria', 'Fecha Adquisición', 'Estado', 'Area'])
        hoja.append(['NUEVA-1', 'Equipo de prueba', '2024-01-01', 'Activo', 'Planta'])
        hoja = libro.create_sheet('Mantenimientos')
        hoja.append(['Código', 'Fecha Mantenimiento', 'Descripción', 'Realizado Por'])
        hoja.append(['NO-EXISTE', '2024-01-01', 'x', 'Nadie'])
        archivo = BytesIO()
        libro.save(archivo)
        archivo.seek(0)
        archivo.name = 'con_error.xlsx'

        respuesta = self.client.post(reverse('upload_csv'), {'file': archivo}, follow=True)
        self.assertContains(respuesta, 'NO-EXISTE')
        self.assertFalse(Maquinaria.objects.filter(codigo_maquinaria='NUEVA-1').exists())

    def test_excel_de_ejemplo_tiene_fechas_de_hoy(self):
        respuesta = self.client.get(reverse('descargar_excel_ejemplo'))
        libro = load_workbook(BytesIO(respuesta.content))
        self.assertIn('Trabajadores', libro.sheetnames)
        self.assertIn(timezone.localdate().strftime('%d-%m-%Y'), libro['Instrucciones']['A2'].value)


class TrabajadoresTests(BaseDemoTest):
    def test_agregar_trabajador_con_certificacion(self):
        juan = Trabajador.objects.get(rut='12345678-9')
        certificacion = Capacitacion.objects.get(nombre_capacitacion='Primeros Auxilios')
        self.client.post(reverse('add_trabajador'), {
            'rut': '9876543-K', 'nombre': 'Persona Nueva', 'area': juan.area_id, 'cargo': juan.cargo_id,
            'turno': juan.turno_id, 'jornada': juan.jornada_id, 'horario': juan.horario_id, 'horas_esperadas': 45,
            'certificacion_nueva_1': certificacion.id, 'fecha_inicio_nueva_1': '2026-01-01', 'fecha_fin_nueva_1': '2027-01-01',
        })
        nuevo = Trabajador.objects.get(rut='9876543-K')
        self.assertEqual(nuevo.capacitaciontrabajador_set.count(), 1)

    def test_editar_trabajador_conserva_turno(self):
        carlos = Trabajador.objects.get(rut='17654321-K')
        respuesta = self.client.get(reverse('editar_trabajador', args=[carlos.rut]))
        self.assertContains(respuesta, f'<option value="{carlos.turno_id}" selected>')

    def test_crear_tipo_de_certificacion(self):
        respuesta = self.client.post(reverse('agregar_capacitacion'), {'nombre': 'Rescate Minero', 'es_renovable': True},
                                     content_type='application/json')
        self.assertTrue(respuesta.json()['success'])
        self.assertTrue(Capacitacion.objects.get(nombre_capacitacion='Rescate Minero').es_renovable)

    def test_dar_de_baja_trabajador(self):
        self.client.post(reverse('eliminar_trabajador'), {'rut': '12345678-9'}, content_type='application/json')
        self.assertFalse(Trabajador.objects.get(rut='12345678-9').activo)


class InventarioTests(BaseDemoTest):
    def test_traspaso_suma_stock_en_panol(self):
        origen = ArticuloBodega.objects.get(nombre_articulo='Filtro de Aceite CAT')
        destino = Panol.objects.get(nombre_panol='Pañol Planta')
        self.client.post(reverse('mover_articulo'), {
            'origen': origen.bodega_id, 'articulo': origen.id, 'cantidad': 5, 'destino': destino.id, 'motivo': 'Prueba',
        })
        origen.refresh_from_db()
        self.assertEqual(origen.cantidad, 25)
        self.assertEqual(ArticuloPanol.objects.get(nombre_articulo='Filtro de Aceite CAT', panol=destino).cantidad, 5)
        self.assertTrue(MovimientoArticulo.objects.filter(motivo='Prueba').exists())

    def test_traspaso_no_permite_mas_que_el_stock(self):
        origen = ArticuloBodega.objects.get(nombre_articulo='Lentes de Seguridad')
        destino = Panol.objects.first()
        self.client.post(reverse('mover_articulo'), {
            'origen': origen.bodega_id, 'articulo': origen.id, 'cantidad': 999, 'destino': destino.id,
        })
        origen.refresh_from_db()
        self.assertEqual(origen.cantidad, 8)

    def test_retiro_y_devolucion(self):
        articulo = ArticuloPanol.objects.get(nombre_articulo='Protector Auditivo')
        trabajador = Trabajador.objects.get(rut='21234567-0')
        self.client.post(reverse('retiro_articulo'), {'trabajador': trabajador.rut, 'articulo': articulo.id, 'cantidad': 4})
        articulo.refresh_from_db()
        self.assertEqual(articulo.cantidad, 66)

        retiro = RetiroArticulo.objects.get(trabajador=trabajador, articulo=articulo)
        self.client.post(reverse('devolver_articulo', args=[retiro.id]))
        articulo.refresh_from_db()
        retiro.refresh_from_db()
        self.assertEqual(articulo.cantidad, 70)
        self.assertTrue(retiro.es_devuelto)


class MaquinariaTests(BaseDemoTest):
    def test_registrar_mantenimiento_deja_equipo_en_mantenimiento(self):
        maquina = Maquinaria.objects.get(codigo_maquinaria='PALA-001')
        self.client.post(reverse('add_mantenimiento'), {
            'maquinaria': maquina.id, 'trabajador': '15876543-2', 'descripcion': 'Cambio de cables',
        })
        maquina.refresh_from_db()
        self.assertEqual(maquina.estado, 'En Mantenimiento')
        self.assertTrue(MantenimientoMaquinaria.objects.filter(descripcion='Cambio de cables').exists())

    def test_registrar_trabajo_calcula_horas(self):
        maquina = Maquinaria.objects.get(codigo_maquinaria='GRUA-001')
        self.client.post(reverse('add_trabajo_maquinaria'), {
            'maquinaria': maquina.id, 'trabajador': '19456789-1', 'horas_trabajadas': 6, 'descripcion': 'Descarga',
        })
        trabajo = TrabajoMaquinaria.objects.get(descripcion='Descarga')
        self.assertEqual(float(trabajo.horas_trabajadas), 6.0)
        self.assertLess(timezone.now() - trabajo.fecha_fin, timedelta(minutes=1))

    def test_cambiar_estado(self):
        maquina = Maquinaria.objects.get(codigo_maquinaria='BULL-001')
        self.client.post(reverse('edit_maquinaria', args=[maquina.id]), {'estado': 'Activo'})
        maquina.refresh_from_db()
        self.assertEqual(maquina.estado, 'Activo')
