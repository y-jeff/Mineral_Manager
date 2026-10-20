"""
Genera datos_prueba/MineralManager_Datos_Prueba.xlsx con datos de ejemplo para probar
todas las secciones de Mineral Manager.

Las fechas se calculan a partir del día de hoy, así siempre hay certificaciones
expiradas, por expirar y vigentes. Uso:

    python datos_prueba/generar_excel.py
"""
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

HOY = date.today()
SALIDA = Path(__file__).resolve().parent / "MineralManager_Datos_Prueba.xlsx"


def dias(n):
    return HOY + timedelta(days=n)


def hora(n_dias, h, m=0):
    d = dias(n_dias)
    return datetime(d.year, d.month, d.day, h, m)


# Columnas: RUT, Nombre, Area, Cargo, Jornada, Turno, Horario, Horas, Horas Esperadas
TRABAJADORES = [
    ("12345678-9", "Juan Pérez Soto", "Extracción", "Operador de Camión", "Diurna", "A", "7x7", 84, 84),
    ("15876543-2", "Ana González Rojas", "Mantenimiento", "Técnico Mecánico", "Diurna", "A", "7x7", 80, 84),
    ("17654321-K", "Carlos Muñoz Díaz", "Extracción", "Operador de Pala", "Nocturna", "B", "7x7", 84, 84),
    ("16234567-8", "María Fernández Vega", "Seguridad", "Prevencionista de Riesgos", "Diurna", "A", "5x2", 45, 45),
    ("18345678-5", "Pedro Rodríguez Lagos", "Planta", "Operador de Planta", "Nocturna", "B", "4x3", 40, 48),
    ("19456789-1", "Camila Torres Fuentes", "Logística", "Encargada de Bodega", "Diurna", "A", "5x2", 45, 45),
    ("14567890-3", "Luis Herrera Pino", "Mantenimiento", "Supervisor de Mantenimiento", "Diurna", "B", "7x7", 84, 84),
    ("20123456-7", "Valentina Rojas Mella", "Planta", "Analista de Laboratorio", "Diurna", "A", "5x2", 38, 45),
    ("13678901-4", "Jorge Castillo Araya", "Extracción", "Perforista", "Nocturna", "B", "7x7", 84, 84),
    ("21234567-0", "Francisca Silva Núñez", "Logística", "Pañolera", "Diurna", "A", "5x2", 45, 45),
    ("11789012-6", "Diego Morales Contreras", "Extracción", "Jefe de Turno", "Diurna", "A", "7x7", 84, 84),
    ("22345678-2", "Javiera Reyes Campos", "Seguridad", "Paramédico", "Nocturna", "B", "4x3", 48, 48),
]

# Columnas: Código, Nombre Maquinaria, Fecha Adquisición, Estado, Area
MAQUINARIAS = [
    ("CAEX-001", "Camión de Extracción CAT 797F", dias(-1800), "Activo", "Extracción"),
    ("CAEX-002", "Camión de Extracción Komatsu 930E", dias(-1500), "En Mantenimiento", "Extracción"),
    ("PALA-001", "Pala Hidráulica Komatsu PC8000", dias(-2200), "Activo", "Extracción"),
    ("PERF-001", "Perforadora Atlas Copco Pit Viper", dias(-1100), "Activo", "Extracción"),
    ("BULL-001", "Bulldozer CAT D11", dias(-2600), "Inactivo", "Extracción"),
    ("CARG-001", "Cargador Frontal CAT 994K", dias(-900), "Activo", "Planta"),
    ("CHAN-001", "Chancador Primario Metso", dias(-3000), "En Mantenimiento", "Planta"),
    ("GRUA-001", "Grúa Horquilla Toyota 3T", dias(-600), "Activo", "Logística"),
]

# Columnas: Bodega, Nombre Articulo, Descripción, Cantidad
BODEGA = [
    ("Bodega Central", "Casco de Seguridad", "Casco clase B con barbiquejo", 120),
    ("Bodega Central", "Guantes de Cuero", "Par de guantes de cabritilla", 45),
    ("Bodega Central", "Lentes de Seguridad", "Lentes claros antiempañante", 8),
    ("Bodega Central", "Filtro de Aceite CAT", "Filtro 1R-0716 para motores CAT", 30),
    ("Bodega Norte", "Neumático 59/80R63", "Neumático para camión de extracción", 4),
    ("Bodega Norte", "Aceite Hidráulico 15W40", "Tambor de 208 litros", 0),
    ("Bodega Norte", "Mangueras Hidráulicas", "Manguera 1\" alta presión, tramo 3 m", 60),
    ("Bodega Sur", "Explosivo ANFO", "Saco 25 kg (acceso restringido)", 200),
    ("Bodega Sur", "Detonadores Electrónicos", "Caja de 50 unidades", 15),
]

# Columnas: Pañol, Nombre Articulo, Descripción, Cantidad
PANOL = [
    ("Pañol Mina", "Llave de Torque", "Llave de torque 1/2\" 40-200 Nm", 6),
    ("Pañol Mina", "Taladro Inalámbrico", "Taladro 18V con 2 baterías", 12),
    ("Pañol Mina", "Linterna Minera", "Lámpara LED para casco", 35),
    ("Pañol Mina", "Arnés de Seguridad", "Arnés cuerpo completo 4 argollas", 3),
    ("Pañol Planta", "Multímetro Fluke", "Multímetro digital True RMS", 5),
    ("Pañol Planta", "Juego de Llaves Allen", "Set métrico 1.5 a 10 mm", 18),
    ("Pañol Planta", "Esmeril Angular", "Esmeril 4 1/2\" 850 W", 0),
    ("Pañol Planta", "Protector Auditivo", "Fono tipo copa 27 dB", 70),
]

# Columnas: RUT, Capacitacion, Fecha Inicio, Fecha Fin (vacía = no renovable)
CAPACITACIONES = [
    ("12345678-9", "Licencia Interna de Conducción", dias(-700), dias(-20)),   # expirada
    ("12345678-9", "Inducción Hombre Nuevo", dias(-900), None),
    ("15876543-2", "Trabajo en Altura", dias(-340), dias(25)),                  # expira < 30 días
    ("15876543-2", "Inducción Hombre Nuevo", dias(-800), None),
    ("17654321-K", "Licencia Interna de Conducción", dias(-300), dias(65)),     # expira < 90 días
    ("16234567-8", "Primeros Auxilios", dias(-200), dias(530)),                 # vigente
    ("18345678-5", "Manejo de Sustancias Peligrosas", dias(-400), dias(-5)),    # expirada
    ("19456789-1", "Operación de Grúa Horquilla", dias(-100), dias(630)),       # vigente
    ("14567890-3", "Trabajo en Altura", dias(-330), dias(80)),                  # expira < 90 días
    ("14567890-3", "Bloqueo y Etiquetado (LOTO)", dias(-150), dias(215)),       # vigente
    ("13678901-4", "Manipulación de Explosivos", dias(-350), dias(12)),         # expira < 30 días
    ("21234567-0", "Inducción Hombre Nuevo", dias(-60), None),
    ("11789012-6", "Primeros Auxilios", dias(-500), dias(230)),                 # vigente
    ("22345678-2", "Primeros Auxilios", dias(-720), dias(-1)),                  # expirada
]

# Columnas: Código, Fecha Mantenimiento, Descripción, Realizado Por (nombre exacto del trabajador)
MANTENIMIENTOS = [
    ("CAEX-001", dias(-90), "Cambio de aceite y filtros de motor", "Ana González Rojas"),
    ("CAEX-002", dias(-3), "Revisión de sistema hidráulico por fuga", "Luis Herrera Pino"),
    ("PALA-001", dias(-45), "Cambio de dientes del balde", "Ana González Rojas"),
    ("CHAN-001", dias(-1), "Reemplazo de revestimiento del manto", "Luis Herrera Pino"),
    ("PERF-001", dias(-30), "Inspección de barras y bit", "Ana González Rojas"),
]

# Columnas: Código, RUT, Fecha Inicio, Fecha Fin, Descripción
TRABAJOS = [
    ("CAEX-001", "12345678-9", hora(-2, 8), hora(-2, 20), "Transporte de mineral fase 4"),
    ("CAEX-001", "12345678-9", hora(-1, 8), hora(-1, 19, 30), "Transporte de estéril a botadero"),
    ("PALA-001", "17654321-K", hora(-2, 20), hora(-1, 8), "Carguío frente norte"),
    ("PERF-001", "13678901-4", hora(-3, 21), hora(-2, 7), "Perforación malla de tronadura 12"),
    ("CARG-001", "18345678-5", hora(-1, 9), hora(-1, 17), "Alimentación de chancador"),
    ("GRUA-001", "19456789-1", hora(-1, 10), hora(-1, 13), "Recepción de neumáticos en bodega"),
]

HOJAS = [
    ("Trabajadores", ["RUT", "Nombre", "Area", "Cargo", "Jornada", "Turno", "Horario", "Horas", "Horas Esperadas"], TRABAJADORES),
    ("Maquinarias", ["Código", "Nombre Maquinaria", "Fecha Adquisición", "Estado", "Area"], MAQUINARIAS),
    ("Inventario Bodega", ["Bodega", "Nombre Articulo", "Descripción", "Cantidad"], BODEGA),
    ("Inventario Pañol", ["Pañol", "Nombre Articulo", "Descripción", "Cantidad"], PANOL),
    ("Capacitaciones", ["RUT", "Capacitacion", "Fecha Inicio", "Fecha Fin"], CAPACITACIONES),
    ("Mantenimientos", ["Código", "Fecha Mantenimiento", "Descripción", "Realizado Por"], MANTENIMIENTOS),
    ("Trabajos Maquinaria", ["Código", "RUT", "Fecha Inicio", "Fecha Fin", "Descripción"], TRABAJOS),
]

INSTRUCCIONES = [
    "Mineral Manager — Datos de prueba",
    f"Generado el {HOY:%d-%m-%Y}. Las fechas son relativas a ese día.",
    "",
    "Cómo usarlo:",
    "1. Inicia sesión y entra a «Subir Archivo».",
    "2. Sube este mismo .xlsx: cada hoja se carga como su tabla, en el orden de las pestañas.",
    "3. Si una hoja tiene un error, no se guarda nada del archivo.",
    "",
    "Reglas:",
    "• No cambies los nombres de las columnas: así se reconoce cada tabla.",
    "• El orden importa: Trabajadores y Maquinarias van antes de Capacitaciones, Mantenimientos y Trabajos.",
    "• Turno: A o B. Estado de maquinaria: Activo, En Mantenimiento o Inactivo.",
    "• Capacitaciones: sin «Fecha Fin» = no renovable.",
    "• Mantenimientos: «Realizado Por» debe ser el nombre exacto de un trabajador.",
    "• Trabajos: fechas con hora (AAAA-MM-DD HH:MM); las horas trabajadas se calculan solas.",
    "• Mantenimientos y Trabajos se agregan cada vez que subes el archivo (no se actualizan).",
    "",
    "Qué deberías ver después de subirlo:",
    "• Inicio: 12 trabajadores, stock bajo en bodega y pañol, 2 máquinas en mantenimiento y 1 inactiva.",
    "• Certificaciones expiradas (3), por expirar en < 30 días (2) y en < 90 días (2).",
    "• Artículos agotados (cantidad 0) y con stock bajo (≤ 10) en Bodega y Pañol.",
    "",
    "Para regenerar con fechas actualizadas: python datos_prueba/generar_excel.py",
]

ENCABEZADO = PatternFill("solid", fgColor="1F2937")


def agregar_hoja(libro, titulo, columnas, filas):
    hoja = libro.create_sheet(titulo)
    hoja.append(columnas)
    for fila in filas:
        hoja.append(list(fila))

    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = ENCABEZADO
        celda.alignment = Alignment(vertical="center")

    for i, col in enumerate(columnas, start=1):
        letra = get_column_letter(i)
        ancho = max(len(str(col)), *(len(str(f[i - 1] or "")) for f in filas))
        hoja.column_dimensions[letra].width = min(ancho + 3, 50)
        for celda in hoja[letra][1:]:
            if isinstance(celda.value, datetime):
                celda.number_format = "yyyy-mm-dd hh:mm"
            elif isinstance(celda.value, date):
                celda.number_format = "yyyy-mm-dd"

    rango = f"A1:{get_column_letter(len(columnas))}{len(filas) + 1}"
    tabla = Table(displayName=titulo.replace(" ", "_").replace("ñ", "n"), ref=rango)
    tabla.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    hoja.add_table(tabla)
    hoja.freeze_panes = "A2"


def main():
    libro = Workbook()
    libro.remove(libro.active)
    for titulo, columnas, filas in HOJAS:
        agregar_hoja(libro, titulo, columnas, filas)

    hoja = libro.create_sheet("Instrucciones")
    for linea in INSTRUCCIONES:
        hoja.append([linea])
    hoja["A1"].font = Font(bold=True, size=14)
    hoja.column_dimensions["A"].width = 110

    libro.save(SALIDA)
    print(f"Excel generado: {SALIDA}")


if __name__ == "__main__":
    main()
